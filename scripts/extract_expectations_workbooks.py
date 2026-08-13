"""Publica los libros de expectativas de Banxico como JSON web."""
from __future__ import annotations
import json, math
from datetime import date, datetime
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
FORECAST_BOOK=ROOT/"data/master/banxico_expectativas_pronosticos.xlsm"
OPINION_BOOK=ROOT/"data/master/banxico_expectativas_monitor_opiniones.xlsm"
OUT=ROOT/"public/data/expectativas-banxico.json"
REPORT=ROOT/"data/expectations-validation-report.json"
def clean(v):
    if v is None or isinstance(v,float) and math.isnan(v): return None
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    return v
def rows(book,sheet="DATA_RAW",header_row=5):
    wb=openpyxl.load_workbook(book,read_only=True,data_only=True,keep_vba=True); ws=wb[sheet]
    headers=[str(c.value or "") for c in ws[header_row]]; out=[]
    for row in ws.iter_rows(min_row=header_row+1,values_only=True):
        if any(v is not None for v in row): out.append({headers[i]:clean(v) for i,v in enumerate(row) if i<len(headers) and headers[i]})
    return out
forecast_fields=["survey_date","variable","variable_group","forecast_type","target_period","target_date","horizon_months","horizon_label","mean","median","q1","q3","min","max","std_dev","n_responses","IQR","Revision_1M","Dispersion_Change","actual_value","forecast_error","indicator_id","period_id"]
opinion_fields=["survey_date","variable","forecast_type","target_period","mean","indicator_id","period_id"]
actual_rows=rows(FORECAST_BOOK,"ACTUALS_RAW",1)
actuals={}
for r in actual_rows:
    key=(r.get("indicator_id"),r.get("period_id"))
    if key[0] and key[1] and r.get("actual_value") is not None: actuals[key]=r.get("actual_value")
forecasts=[]
for raw in rows(FORECAST_BOOK):
    if not raw.get("indicator_id"): continue
    r={k:raw.get(k) for k in forecast_fields}; actual=actuals.get((r["indicator_id"],r["period_id"]))
    r["actual_value"]=actual
    if actual is not None and r["median"] is not None: r["forecast_error"]=r["median"]-actual
    forecasts.append(r)
opinions=[{k:r.get(k) for k in opinion_fields} for r in rows(OPINION_BOOK) if r.get("indicator_id")]
all_rows=forecasts+opinions
payload={"metadata":{"generatedAt":datetime.now().astimezone().isoformat(),"forecastSource":FORECAST_BOOK.name,"opinionSource":OPINION_BOOK.name,"coverage":{"from":min(r["survey_date"] for r in all_rows),"to":max(r["survey_date"] for r in all_rows)}},"config":{"forecastIndicators":sorted({r["indicator_id"] for r in forecasts}),"opinionIndicators":sorted({r["indicator_id"] for r in opinions})},"forecastColumns":forecast_fields,"opinionColumns":opinion_fields,"forecasts":[[r[k] for k in forecast_fields] for r in forecasts],"opinions":[[r[k] for k in opinion_fields] for r in opinions]}
report={"forecasts":len(forecasts),"opinions":len(opinions),"forecastIndicators":payload["config"]["forecastIndicators"],"opinionIndicators":payload["config"]["opinionIndicators"],"actualSeries":len(actuals),"forecastsWithActual":sum(r["actual_value"] is not None for r in forecasts),"forecastsWithError":sum(r["forecast_error"] is not None for r in forecasts),"nullSurveyDates":sum(not r["survey_date"] for r in all_rows)}
OUT.write_text(json.dumps(payload,ensure_ascii=False,separators=(",",":")),encoding="utf-8"); REPORT.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8"); print(json.dumps(report,ensure_ascii=False,indent=2))
if report["nullSurveyDates"]: raise SystemExit(2)
